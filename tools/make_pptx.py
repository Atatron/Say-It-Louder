"""Build deck/You_Yourself_Pitch.pptx from deck/slides.json (same content as deck/index.html)."""
import json, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = os.path.join(os.path.dirname(__file__), "..", "deck")
slides = json.load(open(os.path.join(ROOT, "slides.json"), encoding="utf-8"))
BG, BG2, INK, MUTE, GOLD = (RGBColor(*c) for c in [(0x14,0x14,0x2b),(0x1d,0x1d,0x3a),(0xf6,0xf4,0xee),(0xaa,0xa8,0xc8),(0xf5,0xc0,0x4e)])
SERIF, SANS = "Georgia", "Calibri"
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

def text(sl, x, y, w, h, t, size=20, color=INK, bold=False, font=SANS, align=None):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    paras = t if isinstance(t, list) else [t]
    for k, line in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line; f = r.font; f.size = Pt(size); f.bold = bold; f.color.rgb = color; f.name = font
        p.space_after = Pt(8)
        if align: p.alignment = align
    return tb

def card(sl, x, y, w, h):
    s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = BG2; s.line.color.rgb = RGBColor(0x35,0x35,0x5f); s.adjustments[0] = 0.06; return s

def pic(sl, path, x, y, w=None, h=None):
    return sl.shapes.add_picture(os.path.join(ROOT, path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)

for s in slides:
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.background.fill.solid(); sl.background.fill.fore_color.rgb = BG
    L = s["layout"]
    text(sl, .75, .5, 11.5, .4, s.get("kicker", "").upper(), 12, GOLD, True)
    if L == "title":
        text(sl, .75, 1.6, 11.8, 1.6, s["title"], 72 if len(s["title"]) < 16 else 50, INK, True, SERIF)
        text(sl, .75, 3.4, 11, .6, s.get("sub", ""), 24, GOLD)
        text(sl, .75, 4.5, 11, 1.6, s["lines"], 26, INK, False, SERIF)
    else:
        text(sl, .75, .95, 11.8, 1.2, s["title"], (52 if L == "statement" else 34 if len(s["title"]) <= 42 else 28), INK, True, SERIF)
    if L == "story":
        y = 2.2
        if s.get("quote"):
            text(sl, .75, y, 11, 1, "“" + s["quote"] + "”", 44, GOLD, True, SERIF); y += 1.2
        text(sl, .75, y, 10.6, 4, s["lines"], 20)
    elif L == "statement":
        text(sl, .75, 2.6, 10.6, 4, s["lines"], 24)
    elif L == "drawing":
        for k, im in enumerate(s["images"]):
            x = .75 + k * 4.1
            pic(sl, im, x, 1.7, 3.5); text(sl, x, 4.4, 3.8, 1.2, s["captions"][k], 11.5, RGBColor(0xd6,0xd4,0xea))
        cd = CategoryChartData(); cd.categories = s["chart"]["labels"]; cd.add_series(s["chart"]["title"], s["chart"]["values"])
        gf = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.75), Inches(5.45), Inches(6), Inches(1.5), cd); ch = gf.chart
        ch.has_legend = False; ch.has_title = False; pl = ch.plots[0]; pl.has_data_labels = True
        pl.data_labels.font.size = Pt(12); pl.data_labels.font.color.rgb = INK; pl.gap_width = 60
        pl.series[0].format.fill.solid(); pl.series[0].format.fill.fore_color.rgb = GOLD
        ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False
        ch.category_axis.tick_labels.font.size = Pt(11); ch.category_axis.tick_labels.font.color.rgb = MUTE
        text(sl, 7, 5.9, 5.6, .8, s["chart"]["title"], 13, MUTE)
    elif L == "stats":
        for k, (big, lab) in enumerate(s["stats"]):
            x = .75 + k * 4.1; card(sl, x, 2.2, 3.8, 2.3)
            text(sl, x + .2, 2.3, 3.4, .9, big, 38, GOLD, True, SERIF); text(sl, x + .2, 3.2, 3.4, 1.3, lab, 12.5, RGBColor(0xcf,0xcd,0xe6))
        text(sl, .75, 4.75, 11.8, 2, s["lines"], 17)
    elif L == "image":
        text(sl, .75, 2.2, 4.8, 4, s["lines"], 17)
        pic(sl, s["image"], 5.9, 2.2, 6.7)
        if s.get("cta"):
            url = "you-yourself.vercel.app/" + ("ide.html" if s["id"] == "hardware" else "#sim")
            text(sl, .75, 6.3, 4.8, .5, s["cta"] + " → " + url, 13, GOLD, True)
    elif L == "list":
        for k, (a, b) in enumerate(s["items"]):
            y = 2.15 + k * .95; card(sl, .75, y, 11.8, .82)
            text(sl, .95, y + .12, 3, .6, a, 19, GOLD, True, SERIF); text(sl, 4, y + .14, 8.4, .6, b, 15)
    elif L == "sources":
        text(sl, .75, 1.8, 11.8, 5.4, [f"{k+1}. {t}" for k, t in enumerate(s["items"])], 11.5, RGBColor(0xe4,0xe2,0xf2))
    if s.get("source"): text(sl, .75, 6.95, 11.8, .4, s["source"], 10, MUTE)
    if s.get("notes"): sl.notes_slide.notes_text_frame.text = s["notes"]

out = os.path.join(ROOT, "You_Yourself_Pitch.pptx"); prs.save(out); print("saved", out, len(slides), "slides")
